from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def export_report(file_path, conversion=None, interest=None):
    """Создаёт Word-отчёт по последнему переводу и расчёту процентов"""
    document = Document()
    document.add_heading("Отчёт по валютным операциям", 0)
    document.add_paragraph("Лабораторная работа №2 — Курсы валют")

    if conversion is not None:
        source = conversion["source"]
        target = conversion["target"]
        is_transfer = "source_account_id" in conversion and "target_account_id" in conversion
        document.add_heading(
            "Перевод между счетами" if is_transfer else "Конвертация валют", level=1
        )

        if "transfer_id" in conversion:
            document.add_paragraph(f"Номер операции: {conversion['transfer_id']}")
        if is_transfer:
            document.add_paragraph(f"Счёт отправителя: №{conversion['source_account_id']} ({source})")
            document.add_paragraph(f"Счёт получателя: №{conversion['target_account_id']} ({target})")

        # Основные данные перевода
        lines = [
            f"Сумма списания: {conversion['amount']:.2f} {source}",
            f"Обменный курс: 1 {source} = {conversion['exchange_rate']:.4f} {target}",
            f"Сумма после обмена: {conversion['converted_amount']:.2f} {target}",
            f"Комиссия ({conversion['commission_percent']}%): {conversion['commission']:.2f} {target}",
            f"Итого к зачислению: {conversion['final_amount']:.2f} {target}"
        ]
        for line in lines:
            document.add_paragraph(line)

        # Показываем остатки только тогда, когда они есть в данных
        if "source_balance" in conversion:
            document.add_paragraph(
                f"Остаток счёта отправителя: {conversion['source_balance']:.2f} {source}"
            )
        if "target_balance" in conversion:
            document.add_paragraph(
                f"Баланс счёта получателя: {conversion['target_balance']:.2f} {target}"
            )

    if interest is not None:
        currency = interest["currency"]
        document.add_heading("Расчёт процентов по счёту", level=1)
        for line in [
            f"Баланс для расчёта: {interest['balance']:.2f} {currency}",
            f"Годовая ставка: {interest['annual_rate']}%",
            f"Количество дней: {interest['days']}",
            f"Рассчитанные проценты: {interest['result']:.2f} {currency}",
        ]:
            document.add_paragraph(line)

    document.save(file_path)


def export_excel(filename, table_data):
    """Сохраняет последние записи каждой таблицы SQLite в Excel"""
    sheets = {
        "accounts": "Счета",
        "deposits": "Пополнения",
        "transfers": "Переводы",
        "interest_calculations": "Проценты"
    }
    money_columns = {
        "balance", "amount", "converted_amount", "commission",
        "received_amount", "interest_amount"
    }

    workbook = Workbook()
    workbook.remove(workbook.active)  # Убираем стандартный пустой лист

    for table_name, title in sheets.items():
        worksheet = workbook.create_sheet(title)
        columns, rows = table_data[table_name]
        worksheet.append(columns)
        for row in rows:
            worksheet.append(row)

        # Выделяем шапку и закрепляем её при прокрутке
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="383838")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        # Подбираем простой формат для каждого столбца
        for index, column in enumerate(columns, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = 22
            if column in money_columns:
                number_format = "#,##0.00"
            elif column == "exchange_rate":
                number_format = "0.000000"
            elif column == "commission_percent":
                number_format = "0.##"
            else:
                continue

            for cells in worksheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell = cells[0]
                if isinstance(cell.value, (int, float)):
                    cell.number_format = number_format

    workbook.save(filename)
