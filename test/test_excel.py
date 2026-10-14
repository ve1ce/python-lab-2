import pytest
from openpyxl import load_workbook

import database
from currency_package.reports import export_excel


TABLES = ("accounts", "deposits", "transfers", "interest_calculations")


@pytest.fixture
def test_db(tmp_path, monkeypatch):
    """Каждый тест работает со своей временной базой."""
    monkeypatch.setattr(database, "DB_PATH", tmp_path / "test_excel.db")
    database.initialize_database()


def test_excel_export(tmp_path, test_db):
    """Все четыре листа содержат сохранённые операции."""
    usd = database.create_account("USD")
    eur = database.create_account("EUR")
    database.deposit_to_account(usd.number, 500)
    database.execute_transfer(usd.number, eur.number, 100, 2)
    database.save_interest({
        "currency": "USD", "balance": 500,
        "annual_rate": 5, "days": 30, "result": 2.05
    })

    data = {name: database.get_table_data(name) for name in TABLES}
    file_path = tmp_path / "report.xlsx"
    export_excel(file_path, data)

    assert file_path.exists()
    workbook = load_workbook(file_path, data_only=True)
    try:
        assert workbook.sheetnames == ["Счета", "Пополнения", "Переводы", "Проценты"]
        assert workbook["Счета"]["C2"].value == pytest.approx(88.2)
        assert workbook["Пополнения"]["C2"].value == 500
        assert workbook["Переводы"]["L2"].value == pytest.approx(88.2)
        assert workbook["Проценты"]["G2"].value == pytest.approx(2.05)
    finally:
        workbook.close()


def test_excel_limit(tmp_path, test_db):
    """В Excel попадают только последние 20 счетов."""
    for _ in range(25):
        database.create_account("USD")

    data = {name: database.get_table_data(name) for name in TABLES}
    file_path = tmp_path / "limit.xlsx"
    export_excel(file_path, data)

    workbook = load_workbook(file_path, read_only=True)
    try:
        rows = list(workbook["Счета"].values)
        assert len(rows) == 21  # Заголовок + 20 счетов
        assert rows[1][0] == 25  # Самый новый счёт
        assert rows[-1][0] == 6  # Самый старый из последних 20
    finally:
        workbook.close()
